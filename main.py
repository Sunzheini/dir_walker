import os
import time
from os import walk
from os.path import exists
from tkinter import filedialog

import openpyxl
from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.styles import Font

from tkinter import *


# custom functions
# --------------------------------------------------------------------

def time_measurement_decorator(some_function):
    def wrapper(*args, **kwargs):
        start = time.time()
        result = some_function(*args, **kwargs)     # the function
        end = time.time()
        measurement = end - start
        label3.config(text=f"generated in: {measurement:.5f}")
        return result                               # the function
    return wrapper


def empty_the_output_txt_file(file_path):
    if exists(file_path):
        os.remove(file_path)


def empty_the_output_excel_file(file_path):
    if exists(file_path):
        os.remove(file_path)


def write_into_txt_file(file_path, content):
    with open(file_path, 'a') as output_file:
        output_file.write(content + '\n')


def write_into_excel_file(file_path, content):
    workbook = Workbook()                 # creates new workbook
    worksheet = workbook.active
    worksheet.title = "Data"

    # worksheet.append(['Daniel', 'is', 'superman'])  # add to the end of the content, first row on diff cols
    # worksheet.append(['Maxi', 'is', 'superman'])    # next row on different columns

    global current_excel_cell
    worksheet[current_excel_cell] = content

    number = int(current_excel_cell[1])
    number += 1

    current_excel_cell = current_excel_cell[0] + str(number)

    workbook.save(file_path)


def open_dir():
    filepath = filedialog.askdirectory()
    global current_folder
    current_folder = filepath
    label2.config(text=f"selected: '{current_folder}'")


def fix_the_dir_path_str(dir_path):
    dir_path = dir_path[0:2] + '\\' + dir_path[3:]
    return dir_path


@time_measurement_decorator
def the_walk_loop():
    empty_the_output_txt_file(generated_txt_file_path)
    empty_the_output_excel_file(generated_txt_file_path)

    folder = current_folder

    for (dir_path, dir_names, file_names) in walk(folder):

        dir_path = fix_the_dir_path_str(dir_path)

        write_into_txt_file(generated_txt_file_path, f'Directory: {dir_path}')
        for directory in dir_names:
            write_into_txt_file(generated_txt_file_path, f'a directory: {directory}')
            write_into_excel_file(generated_excel_file_path, f'a directory: {directory}')

        for file in file_names:
            write_into_txt_file(generated_txt_file_path, f'a file: {file}')
            write_into_excel_file(generated_excel_file_path, f'a file: {file}')

        write_into_txt_file(generated_txt_file_path, '\n')


# target directory
# --------------------------------------------------------------------
current_folder = ''


# the output files
# --------------------------------------------------------------------
generated_txt_file_path = './generated_txt_file.txt'
generated_excel_file_path = './generated_excel_file.xlsx'
# empty_the_output_txt_file(generated_txt_file_path)

current_excel_cell = 'A1'


# tkinter gui
# --------------------------------------------------------------------

window = Tk()
window.geometry("300x300")
window.title("dir walker")
window.config(background='#2b2828')

label1 = Label(window, text='select target dir', width=30, height=1,
               bg='#2b2828', borderwidth=0, relief="ridge", fg='white')
label1.grid(row=0, column=0)

button1 = Button(window, text='browse', width=30, height=1, command=open_dir)
button1.grid(row=1, column=0)

label2 = Label(window, text=current_folder, width=30, height=1,
               bg='#2b2828', borderwidth=0, relief="ridge", fg='white')
label2.grid(row=2, column=0)

button2 = Button(window, text='generate', width=30, height=1, command=the_walk_loop)
button2.grid(row=3, column=0)

label3 = Label(window, text='', width=30, height=1,
               bg='#2b2828', borderwidth=0, relief="ridge", fg='white')
label3.grid(row=4, column=0)


# gui loop
# --------------------------------------------------------------------
window.mainloop()
